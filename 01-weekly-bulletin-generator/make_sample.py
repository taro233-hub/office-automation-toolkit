# -*- coding: utf-8 -*-
"""生成周进展通报的 demo 样张（纯虚构数据，用于离线演示 generate.py 的输入）。"""
import os

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample", "bulletin_2026-10-12.docx")

HEADERS = ["序号", "工作事项", "本周进展", "下周计划", "备注"]
ROWS = [
    ["1", "A栋幕墙深化设计",
     "第三轮深化图纸完成评审，预埋件定位图已发现场",
     "配合现场完成预埋件首层放线复核",
     "设计院要求周二前反馈意见"],
    ["2", "精装样板段施工",
     "样板段隔墙完成80%，机电点位现场确认完毕",
     "样板段封板，启动乳胶漆基层处理",
     ""],
    ["3", "消防检测委托",
     "检测单位比价完成，合同走内部审批流程",
     "签订检测合同，约定检测进场时间",
     "审批预计本周五前出结果"],
    ["4", "电梯安装调试",
     "3#、4#电梯导轨安装完成，厅门待货",
     "跟进厅门到货，安排轿厢拼装",
     "到货时间供应商口头承诺周三，未书面确认"],
    ["5", "景观铺装样板",
     "铺装材料小样已封样，编号存档",
     "现场做1:1铺装样板，邀请相关方确认",
     ""],
]


def make_sample():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("星河产业园A栋项目周进展通报（2026.10.12–10.16）")
    run.bold = True
    run.font.size = Pt(14)

    meta = doc.add_paragraph("编制日期：2026-10-12　　编制：项目工程组")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1 + len(ROWS), cols=5)
    table.style = "Table Grid"
    for j, h in enumerate(HEADERS):
        cell = table.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for i, row in enumerate(ROWS, start=1):
        for j, v in enumerate(row):
            table.rows[i].cells[j].text = v

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("已生成样张:", OUT)


if __name__ == "__main__":
    make_sample()
