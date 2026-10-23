# -*- coding: utf-8 -*-
"""纸质登记表 OCR 转录器
================================

解决的问题
----------
项目上的用印/验收/巡检类登记表多为纸质手写，每月底要把一摞扫描件逐
行誊进电子台账：人工转录慢、数字和日期容易看错，且事后无法追溯。

流水线
------
扫描件 → PaddleOCR 识别 → 按表头坐标分列 → 规则校验
      → 回写 Excel 台账（存疑格黄底）→ 输出人工复核清单

用法
----
# 常规：对 sample/sheets 下所有扫描样张跑 OCR 并回写台账
python transcribe.py

# 离线演示（不依赖 PaddleOCR/模型，直接用样张标注走完整流水线）
python transcribe.py --mock

# 指定输入目录与输出目录
python transcribe.py --sheets sample/sheets --out output
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
HEADERS = ["序号", "日期", "文件名称", "份数", "用章类型", "经办人"]
VALID_SEAL_TYPES = {"公章", "合同章", "财务章", "法人章", "其他"}
DATE_PATTERNS = [r"(20\d{2})[.\-/年](\d{1,2})[.\-/月](\d{1,2})"]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(vertical="center", wrap_text=True)


# ---------------------------------------------------------------- OCR 层
def ocr_image(path):
    """返回 [(text, cx, cy), ...]；OCR 失败时停止，避免误用演示标注。"""
    try:
        return _ocr_paddle(path)
    except Exception as e:                       # noqa: BLE001
        raise SystemExit("PaddleOCR 识别失败：%s；演示数据请显式使用 --mock" % e) from e


def _ocr_paddle(path):
    from paddleocr import PaddleOCR                # 延迟导入；--mock 无需 OCR 依赖

    ocr = PaddleOCR(lang="ch", use_doc_orientation_classify=False,
                    use_doc_unwarping=False, use_textline_orientation=False,
                    enable_mkldnn=False)   # 规避部分版本 oneDNN 的 PIR 算子兼容问题
    results = ocr.predict(path)
    items = []
    for res in results:
        if isinstance(res, dict) or hasattr(res, "keys"):
            texts = list(res["rec_texts"])
            polys = list(res["rec_polys"])
            for text, poly in zip(texts, polys):
                xs = [p[0] for p in poly]
                ys = [p[1] for p in poly]
                items.append((str(text), sum(xs) / len(xs), sum(ys) / len(ys)))
        else:                                      # 旧版返回格式兼容
            for box, (text, _score) in res:
                xs = [p[0] for p in box]
                ys = [p[1] for p in box]
                items.append((text, sum(xs) / 4, sum(ys) / 4))
    return items


def ocr_mock(path):
    """从样张标注 JSON 生成带网格坐标的伪 OCR 结果，走同一套解析/校验/回写。

    坐标与 make_sample.py 的排版保持一致：
    左边距 90，列宽 [90,210,470,130,210,210]，表头高 70，行高 115。
    """
    gt_path = os.path.join(os.path.dirname(path), "ground_truth.json")
    if not os.path.exists(gt_path):
        raise SystemExit("mock 模式需要 ground_truth.json（先运行 make_sample.py）")
    with open(gt_path, encoding="utf-8") as f:
        gt = json.load(f)

    col_w = [90, 210, 470, 130, 210, 210]
    edges = [90]
    for w in col_w:
        edges.append(edges[-1] + w)
    centers = [(edges[i] + edges[i + 1]) / 2 for i in range(len(col_w))]

    items = [(h, centers[j], 190 + 70 / 2) for j, h in enumerate(HEADERS)]   # 表头
    y = 190 + 70 + 115 / 2                                                    # 第一条数据行中心
    for row in gt["records"]:
        for j, val in enumerate(row):
            if val:
                items.append((val, centers[j], y))
        y += 115
    return items


# ------------------------------------------------------------ 结构化解析
def parse_records(items):
    """按表头列坐标 + 视觉行聚类把 OCR 文本装配成行记录。

    聚类规则：
    - y 间距 > 55px 视为新视觉行（与样张行高 115px 匹配）；
    - 有“序号”的视觉行开启新记录；
    - 与上一视觉行间距 <= 70px 的无序号行视为单元格内换行，拼接到上一条；
    - 其余无序号行独立成记录（序号留空 → 校验拦截为待复核）。
    """
    header_pos = {t: (cx, cy) for t, cx, cy in items if t in HEADERS}
    missing = [h for h in HEADERS if h not in header_pos]
    if missing:
        raise SystemExit("未识别到完整表头，缺少：%s；请人工复核扫描件" % "、".join(missing))
    header_y = min(cy for _, cy in header_pos.values())
    cols = [(h, header_pos[h][0]) for h in HEADERS if h in header_pos]

    # 1) 视觉行聚类
    data = sorted((it for it in items if it[2] > header_y + 30), key=lambda x: (x[2], x[1]))
    clusters, cur, prev_y = [], [], None
    for it in data:
        if prev_y is not None and it[2] - prev_y > 55:
            clusters.append(cur)
            cur = []
        cur.append(it)
        prev_y = it[2]
    if cur:
        clusters.append(cur)

    # 2) 逐行装配
    records, last = [], None
    for i, cluster in enumerate(clusters):
        assigned = {}
        for text, cx, cy in sorted(cluster, key=lambda x: x[1]):
            col = min(cols, key=lambda hc: abs(hc[1] - cx))[0]
            assigned[col] = (assigned[col] + " " + text) if col in assigned else text
        if "序号" in assigned:
            last = dict.fromkeys(HEADERS, "")
            last.update(assigned)
            records.append(last)
        elif last is not None and cluster[0][2] - clusters[i - 1][0][2] <= 70:
            for k, v in assigned.items():                     # 单元格内换行拼接
                last[k] = (last[k] + v) if last[k] else v
        else:
            last = dict.fromkeys(HEADERS, "")                 # 疑似序号漏识别
            last.update(assigned)
            records.append(last)
    return records


# ------------------------------------------------------------- 规则校验
def validate(row, expect_no=None):
    """返回 (该行问题列表, 每字段问题映射)。"""
    problems, field_flags = [], {}

    def flag(field, msg):
        problems.append(msg)
        field_flags[field] = True

    if expect_no is not None and row["序号"] != str(expect_no):
        flag("序号", "序号应为 %s，实为 %r（断档/跳行）" % (expect_no, row["序号"]))

    date = None
    date_raw = re.sub(r"\s+", "", row["日期"])        # OCR 常在数字间插入空格，先归一
    for pat in DATE_PATTERNS:
        m = re.search(pat, date_raw)
        if m:
            y, mth, d = map(int, m.groups())
            try:
                date = dt.date(y, mth, d)
            except ValueError:
                flag("日期", "日期 %r 不是合法日期" % row["日期"])
            break
    if date is None and "日期" not in field_flags:
        flag("日期", "日期 %r 无法解析" % row["日期"])
    elif date and (date > dt.date.today() or date < dt.date(2000, 1, 1)):
        flag("日期", "日期 %s 超出合理范围" % date)

    if not row["份数"].isdigit() or int(row["份数"] or 0) < 1:
        flag("份数", "份数 %r 不是正整数" % row["份数"])
    if not row["文件名称"].strip():
        flag("文件名称", "文件名称为空")
    if row["用章类型"] not in VALID_SEAL_TYPES:
        flag("用章类型", "用章类型 %r 不在白名单 %s" % (row["用章类型"], sorted(VALID_SEAL_TYPES)))
    if not row["经办人"].strip():
        flag("经办人", "经办人为空")
    return problems, field_flags


# ------------------------------------------------------------- 回写台账
def write_ledger(all_rows, out_dir, template):
    out_xlsx = os.path.join(out_dir, "ledger.xlsx")
    if os.path.exists(template):
        shutil.copy(template, out_xlsx)
        wb = load_workbook(out_xlsx)
        ws = wb.active
        while ws.max_row > 1:                       # 清掉模板里的历史数据行
            ws.delete_rows(2)
    else:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS + ["转录状态", "复核标记"])

    report = ["# 台账转录复核清单", ""]
    n_ok = n_review = 0
    for sheet_name, rows in all_rows:
        report.append("## %s" % sheet_name)
        for i, row in enumerate(rows, start=1):
            problems, flags = validate(row, expect_no=i)
            r = ws.max_row + 1
            for j, h in enumerate(HEADERS, start=1):
                cell = ws.cell(row=r, column=j, value=row[h])
                cell.alignment = WRAP
                if h in flags:
                    cell.fill = YELLOW
                else:
                    cell.fill = GREEN
            if problems:
                n_review += 1
                ws.cell(row=r, column=len(HEADERS) + 1, value="待复核")
                ws.cell(row=r, column=len(HEADERS) + 2, value="；".join(problems))
                report.append("- 第%d条 【待复核】%s" % (i, "；".join(problems)))
            else:
                n_ok += 1
                ws.cell(row=r, column=len(HEADERS) + 1, value="已转录")
                report.append("- 第%d条 【已转录】%s %s（%s份，%s）" %
                              (i, row["日期"], row["文件名称"], row["份数"], row["用章类型"]))
        report.append("")

    report.insert(2, "**已转录 %d 条，待复核 %d 条（黄底格 = 转录存疑，需人工对照原件确认）**" % (n_ok, n_review))
    wb.save(out_xlsx)

    report_path = os.path.join(out_dir, "review_report.md")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(report))
    return out_xlsx, report_path, n_ok, n_review


def main():
    ap = argparse.ArgumentParser(description="纸质登记表 OCR 转录到电子台账")
    ap.add_argument("--sheets", default=os.path.join(HERE, "sample", "sheets"))
    ap.add_argument("--out", default=os.path.join(HERE, "output"))
    ap.add_argument("--mock", action="store_true", help="不跑 OCR，用样张标注走流水线")
    args = ap.parse_args()

    pngs = sorted(f for f in os.listdir(args.sheets) if f.lower().endswith(".png"))
    if not pngs:
        raise SystemExit("目录中没有扫描件（.png）: %s" % args.sheets)

    os.makedirs(args.out, exist_ok=True)
    all_rows = []
    for name in pngs:
        path = os.path.join(args.sheets, name)
        items = ocr_mock(path) if args.mock else ocr_image(path)
        rows = parse_records(items)
        print("%-28s 识别 %d 条记录" % (name, len(rows)))
        all_rows.append((name, rows))

    template = os.path.join(HERE, "sample", "ledger_template.xlsx")
    xlsx, report, n_ok, n_review = write_ledger(all_rows, args.out, template)
    print("回写台账 : %s" % xlsx)
    print("复核清单 : %s" % report)
    print("结果     : 已转录 %d 条，待复核 %d 条" % (n_ok, n_review))


if __name__ == "__main__":
    main()
